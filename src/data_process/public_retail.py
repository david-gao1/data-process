"""Bounded public snapshot ingestion. No scheduling, credentials or external uploads."""
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import csv
import hashlib
import io
import json
from pathlib import Path
import re
import subprocess
import tempfile
import zipfile

import openpyxl

from .core import ROOT, ASSET_ROOT, digest, durable_write

OFFICIAL_URL = 'https://archive.ics.uci.edu/static/public/352/online+retail.zip'
HEADERS = ['InvoiceNo', 'StockCode', 'Description', 'Quantity', 'InvoiceDate', 'UnitPrice', 'CustomerID', 'Country']
NORMALIZED = ['business_date', 'country', 'line_type', 'amount_pence', 'source_row']
MAX_WORKBOOK = 32 * 1024 * 1024
MAX_XML = 384 * 1024 * 1024


def file_hash(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(65536), b''):
            h.update(block)
    return h.hexdigest()


def source_spec():
    return json.loads((ASSET_ROOT / 'sources/uci-online-retail.json').read_text())


def transfer_https(url, path, limit):
    """System curl supplies trusted TLS and timeouts; Python enforces received-byte bound."""
    with tempfile.TemporaryFile() as errors:
        process = subprocess.Popen(['curl', '--proto', '=https', '--fail', '--silent', '--show-error',
                                    '--max-time', '120', '--connect-timeout', '15', url],
                                   stdout=subprocess.PIPE, stderr=errors)
        count = 0
        try:
            with Path(path).open('wb') as output:
                while True:
                    block = process.stdout.read(65536)
                    if not block:
                        break
                    count += len(block)
                    if count > limit:
                        raise ValueError('download exceeds byte limit')
                    output.write(block)
                output.flush()
                import os
                os.fsync(output.fileno())
            code = process.wait(timeout=5)
            if code != 0:
                raise ValueError(f'HTTPS download failed: curl exit {code}; TLS verification remains enabled')
        finally:
            process.stdout.close()
            if process.poll() is None:
                process.terminate()
                try:
                    process.wait(timeout=2)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.wait()
    return {'received_bytes': count, 'transport': 'system curl; TLS verified; no redirects/retries'}


def inspect_archive(path, spec):
    if not zipfile.is_zipfile(path):
        raise ValueError('invalid ZIP source')
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) != 1 or members[0].filename != spec['xlsx_member']:
            raise ValueError('unexpected archive members')
        if members[0].file_size > MAX_WORKBOOK:
            raise ValueError('workbook exceeds size limit')
        workbook = archive.read(members[0])
    if digest(workbook) != spec['expected_xlsx_sha256']:
        raise ValueError('workbook checksum mismatch')
    with zipfile.ZipFile(io.BytesIO(workbook)) as archive:
        if sum(m.file_size for m in archive.infolist()) > MAX_XML:
            raise ValueError('XLSX expansion exceeds size limit')
    return workbook


def fetch(p, spec=None, transfer=transfer_https):
    spec = spec or source_spec()
    if spec['url'] != OFFICIAL_URL or spec['license'] != 'CC-BY-4.0' or spec['dataset'] != 'uci-online-retail-352':
        raise ValueError('source is not the registered official dataset')
    expected = spec.get('expected_sha256', '')
    if not isinstance(expected, str) or not re.fullmatch('[0-9a-f]{64}', expected):
        raise ValueError('a pinned source SHA256 is required')
    limit = spec['max_bytes']
    if type(limit) is not int or not 1 <= limit <= 32 * 1024 * 1024:
        raise ValueError('invalid source byte limit')
    raw = p.home / 'raw' / f'{expected}.zip'
    if raw.exists():
        if raw.stat().st_size > limit or file_hash(raw) != expected:
            raise ValueError('cached raw checksum mismatch')
        inspect_archive(raw, spec)
        if not raw.with_suffix('.json').exists():
            recovered = dict(spec, raw_sha256=expected, downloaded_bytes=raw.stat().st_size,
                             retrieved_utc='unknown: validated cache recovered after metadata interruption')
            durable_write(raw.with_suffix('.json'), json.dumps(recovered, sort_keys=True).encode())
        return raw
    raw.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='download-', dir=p.home) as temporary:
        temporary_path = Path(temporary) / 'source.zip'
        transfer(spec['url'], temporary_path, limit)
        if temporary_path.stat().st_size > limit or file_hash(temporary_path) != expected:
            raise ValueError('source size or checksum mismatch')
        inspect_archive(temporary_path, spec)
        durable_write(raw, temporary_path.read_bytes())
    manifest = dict(spec, raw_sha256=expected, downloaded_bytes=raw.stat().st_size,
                    retrieved_utc=datetime.utcnow().isoformat() + 'Z')
    durable_write(raw.with_suffix('.json'), json.dumps(manifest, sort_keys=True).encode())
    return raw


def normalized_row(row, ordinal):
    invoice, _stock, _description, quantity, when, price, _customer, country = row
    if not isinstance(when, datetime) or when.tzinfo is not None:
        return None, 'invalid_date', False
    if not isinstance(country, str) or not country.strip() or not invoice:
        return None, 'missing_dimension', False
    if type(quantity) not in (int, float) or not float(quantity).is_integer() or quantity == 0:
        return None, 'invalid_quantity', False
    try:
        money = Decimal(str(price))
    except (InvalidOperation, ValueError):
        return None, 'invalid_price', False
    if not money.is_finite() or money <= 0:
        return None, 'nonpositive_price', False
    cancelled = str(invoice).strip().upper().startswith('C')
    if cancelled != (quantity < 0):
        return None, 'cancellation_sign_mismatch', False
    exact = abs(Decimal(str(quantity)) * money * 100)
    pence = exact.quantize(Decimal('1'), rounding=ROUND_HALF_UP)
    if pence > 10**14:
        return None, 'amount_out_of_range', False
    return [when.date().isoformat(), country.strip(), 'refund' if cancelled else 'sale', int(pence), ordinal], None, exact != pence


def ingest_month(p, month, spec=None, raw_path=None):
    if not re.fullmatch(r'20\d\d-(0[1-9]|1[0-2])', month):
        raise ValueError('batch month must be YYYY-MM')
    spec = spec or source_spec()
    raw = Path(raw_path) if raw_path else fetch(p, spec)
    if file_hash(raw) != spec['expected_sha256']:
        raise ValueError('raw source checksum mismatch')
    workbook = inspect_archive(raw, spec)
    transform_hash = file_hash(Path(__file__))
    batch_id = digest(f'{spec["dataset"]}:{month}:{spec["expected_sha256"]}:{transform_hash}:{openpyxl.__version__}'.encode())
    manifest_path = p.home / 'ingestions' / f'{batch_id}.json'
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text())
        verify_batch(p, manifest)
        return manifest
    book = openpyxl.load_workbook(io.BytesIO(workbook), read_only=True, data_only=False, keep_links=False)
    try:
        if len(book.sheetnames) != 1:
            raise ValueError('expected exactly one source worksheet')
        sheet = book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        if list(next(rows)) != HEADERS:
            raise ValueError('source schema mismatch: exact expected columns are required')
        by_day = defaultdict(list)
        rejects = Counter()
        reject_rows = []
        total = selected = rounded = missing_customer = missing_description = 0
        observed = Counter()
        for ordinal, row in enumerate(rows, start=2):
            total += 1
            if total > 600000:
                raise ValueError('source row limit exceeded')
            if len(row) != 8:
                raise ValueError('source row schema mismatch')
            invoice, _, description, quantity, when, price, customer, _ = row
            missing_customer += customer is None
            missing_description += description is None
            observed['cancellation_rows'] += str(invoice).upper().startswith('C')
            observed['negative_quantity_rows'] += isinstance(quantity, (int, float)) and quantity < 0
            observed['zero_price_rows'] += price == 0
            observed['negative_price_rows'] += isinstance(price, (int, float)) and price < 0
            if not isinstance(when, datetime):
                raise ValueError('unattributable source date; batch cannot safely select a month')
            if when.strftime('%Y-%m') != month:
                continue
            selected += 1
            cleaned, reason, was_rounded = normalized_row(row, ordinal)
            if reason:
                rejects[reason] += 1
                reject_rows.append([ordinal, reason])
                continue
            rounded += was_rounded
            by_day[cleaned[0]].append(cleaned)
    finally:
        book.close()
    rejected = sum(rejects.values())
    if selected == 0 or not by_day or rejected / selected > 0.05:
        raise ValueError('batch has no usable records or exceeds 5% quarantine threshold')
    sources = []
    expected = {'gross_sales_pence': 0, 'refund_pence': 0, 'sale_lines': 0, 'refund_lines': 0}
    for day, records in sorted(by_day.items()):
        output = io.StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(NORMALIZED)
        writer.writerows(records)
        data = output.getvalue().encode()
        source_file = p.home / 'ingestions' / 'normalized' / f'{batch_id}-{day}.csv'
        durable_write(source_file, data)
        source_id = p.ready(day, source_file)
        sources.append({'date': day, 'source_id': source_id, 'normalized_sha256': digest(data), 'rows': len(records)})
        for record in records:
            if record[2] == 'sale':
                expected['gross_sales_pence'] += record[3]
                expected['sale_lines'] += 1
            else:
                expected['refund_pence'] += record[3]
                expected['refund_lines'] += 1
    expected['net_sales_pence'] = expected['gross_sales_pence'] - expected['refund_pence']
    manifest = {'batch_id': batch_id, 'month': month, 'dataset': spec['dataset'], 'source_url': spec['url'],
                'source_page': spec['page'], 'license': spec['license'], 'attribution': spec['attribution'],
                'raw_sha256': spec['expected_sha256'], 'xlsx_sha256': spec['expected_xlsx_sha256'],
                'sheet': sheet.title, 'row_identity': 'raw_sha256 + sheet + source_row (1-based XLSX row)',
                'normalizer_sha256': transform_hash, 'openpyxl': openpyxl.__version__,
                'snapshot_rows': total, 'snapshot_missing_customer': missing_customer, 'snapshot_missing_description': missing_description,
                'snapshot_profile': dict(observed), 'selected_rows': selected, 'quarantined_rows': rejected,
                'quarantine_reasons': dict(rejects), 'quarantine_row_ids': reject_rows, 'rounded_lines': rounded,
                'sources': sources, 'expected_totals': expected, 'incremental_semantics': 'historical monthly replay of one static snapshot; not CDC'}
    manifest['manifest_sha256'] = digest(json.dumps(manifest, ensure_ascii=False, sort_keys=True).encode())
    durable_write(manifest_path, json.dumps(manifest, ensure_ascii=False, sort_keys=True).encode())
    return manifest


def verify_batch(p, manifest):
    unsigned = dict(manifest)
    recorded_hash = unsigned.pop('manifest_sha256')
    if digest(json.dumps(unsigned, ensure_ascii=False, sort_keys=True).encode()) != recorded_hash:
        raise ValueError('batch manifest checksum mismatch')
    raw = p.home / 'raw' / f'{manifest["raw_sha256"]}.zip'
    if file_hash(raw) != manifest['raw_sha256']:
        raise ValueError('batch raw checksum mismatch')
    for source in manifest['sources']:
        path = p.home / 'sources' / f'{source["source_id"]}.csv'
        if file_hash(path) != source['normalized_sha256']:
            raise ValueError('normalized source checksum mismatch')
    return {'batch_id': manifest['batch_id'], 'integrity': 'ok', 'dates': len(manifest['sources'])}


def stage_batch(p, manifest, release_id):
    verify_batch(p, manifest)
    return [p.submit(release_id, s['source_id'], s['date'], timeout=120,
                     key=f'retail:{manifest["batch_id"]}:{s["date"]}:{release_id}') for s in manifest['sources']]


def main():
    import argparse
    from .core import Platform
    parser = argparse.ArgumentParser(description='Public historical retail snapshot; no automatic network schedule')
    parser.add_argument('--home', required=True)
    sub = parser.add_subparsers(dest='action', required=True)
    sub.add_parser('fetch')
    command = sub.add_parser('ingest')
    command.add_argument('--month', required=True)
    for name in ('verify', 'stage'):
        command = sub.add_parser(name)
        command.add_argument('--batch', required=True)
        if name == 'stage':
            command.add_argument('--release', required=True)
    command = sub.add_parser('trace')
    command.add_argument('--source', required=True)
    args = parser.parse_args()
    p = Platform(args.home)
    p.init()
    try:
        if args.action == 'fetch':
            raw = fetch(p)
            result = {'raw_sha256': file_hash(raw), 'bytes': raw.stat().st_size, 'path': str(raw)}
        elif args.action == 'ingest':
            manifest = ingest_month(p, args.month)
            result = {key: value for key, value in manifest.items() if key != 'quarantine_row_ids'}
        elif args.action == 'trace':
            matches = []
            for path in (p.home / 'ingestions').glob('*.json'):
                manifest = json.loads(path.read_text())
                for source in manifest['sources']:
                    if source['source_id'] == args.source:
                        verify_batch(p, manifest)
                        matches.append({k: manifest[k] for k in ('batch_id','raw_sha256','xlsx_sha256','month','source_page','license','attribution','row_identity')})
            if not matches:
                raise ValueError('source has no verified public lineage')
            result = matches
        else:
            if not re.fullmatch('[0-9a-f]{64}', args.batch):
                raise ValueError('invalid batch id')
            manifest = json.loads((p.home / 'ingestions' / f'{args.batch}.json').read_text())
            result = verify_batch(p, manifest) if args.action == 'verify' else {'run_ids': stage_batch(p, manifest, args.release)}
        print(json.dumps(result, ensure_ascii=False, indent=2))
    except (ValueError, OSError, KeyError, zipfile.BadZipFile) as exc:
        parser.exit(2, f'public retail failed: {exc}\n')


if __name__ == '__main__':
    main()
